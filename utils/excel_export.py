"""
Bek Rate Desk — Excel Export Utility
======================================
openpyxl-based export for NII Sensitivity, EVE/Duration Gap,
RV Matrix, and Repricing Gap outputs.

Usage:
  from utils.excel_export import export_nii, export_eve, export_rv_matrix, export_repricing_gap
  buf = export_nii(nii_result)   # returns BytesIO
"""

import io
from datetime import datetime
from typing import Any

try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import (
        PatternFill, Font, Alignment, Border, Side, numbers
    )
    from openpyxl.utils import get_column_letter
    _HAS_OPENPYXL = True
except ImportError:
    _HAS_OPENPYXL = False


# ── Colour palette (dark-theme inspired, works on white paper too) ────────────
C_NAVY      = "1E2233"
C_MIDNIGHT  = "0E1017"
C_GOLD      = "F59E0B"
C_GREEN     = "10B981"
C_RED       = "EF4444"
C_AMBER     = "FBBF24"
C_HEADER_BG = "1E2233"
C_ALT_ROW   = "F8F9FF"
C_WHITE     = "FFFFFF"
C_SUBHEAD   = "2563EB"


def _require_openpyxl():
    if not _HAS_OPENPYXL:
        raise ImportError("openpyxl is required. Install with: pip install openpyxl")


def _header_font(size=10, bold=True, color=C_WHITE):
    return Font(name="Calibri", size=size, bold=bold, color=color)


def _body_font(size=9, bold=False, color="000000"):
    return Font(name="Calibri", size=size, bold=bold, color=color)


def _fill(hex_color):
    return PatternFill(fill_type="solid", fgColor=hex_color)


def _border():
    thin = Side(style="thin", color="D1D5DB")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _center():
    return Alignment(horizontal="center", vertical="center", wrap_text=False)


def _right():
    return Alignment(horizontal="right", vertical="center")


def _set_col_width(ws, col_letter, width):
    ws.column_dimensions[col_letter].width = width


def _stamp_header(ws, title: str, subtitle: str = ""):
    """Writes a 2-row title block at the top of a worksheet."""
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 14
    cell = ws.cell(row=1, column=1, value=title)
    cell.font = Font(name="Calibri", size=13, bold=True, color=C_GOLD)
    cell.fill = _fill(C_NAVY)
    cell.alignment = Alignment(horizontal="left", vertical="center")

    ts = datetime.now().strftime("%d %b %Y  %H:%M")
    ts_cell = ws.cell(row=1, column=2, value=ts)
    ts_cell.font = Font(name="Calibri", size=9, color="94A3B8")
    ts_cell.fill = _fill(C_NAVY)
    ts_cell.alignment = _right()

    if subtitle:
        sub = ws.cell(row=2, column=1, value=subtitle)
        sub.font = Font(name="Calibri", size=9, color="64748B")
        sub.fill = _fill("F1F5F9")


def _write_header_row(ws, row: int, cols: list, bg=C_HEADER_BG, fg=C_WHITE, height=15):
    ws.row_dimensions[row].height = height
    for ci, label in enumerate(cols, start=1):
        c = ws.cell(row=row, column=ci, value=label)
        c.font = _header_font(size=9, color=fg)
        c.fill = _fill(bg)
        c.alignment = _center()
        c.border = _border()


def _write_data_row(ws, row: int, values: list, alt=False, formats=None):
    bg = C_ALT_ROW if alt else C_WHITE
    for ci, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=ci, value=v)
        c.font = _body_font()
        c.fill = _fill(bg)
        c.border = _border()
        c.alignment = _right() if isinstance(v, (int, float)) else Alignment(vertical="center")
        if formats and ci <= len(formats) and formats[ci - 1]:
            c.number_format = formats[ci - 1]


# ── NII SENSITIVITY EXPORT ────────────────────────────────────────────────────

def export_nii(result: dict) -> io.BytesIO:
    """
    Export NII Sensitivity results to Excel.
    result = output of alm_engine.run_nii_sensitivity()
    """
    _require_openpyxl()
    wb = Workbook()

    # ── Sheet 1: Summary ──────────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "NII Summary"
    _stamp_header(ws1, "NII Sensitivity Analysis", "Bek Rate Desk — ALM Module")
    ws1.merge_cells("A1:F1")
    ws1.merge_cells("A2:F2")

    r = 4
    # Key metrics
    _write_header_row(ws1, r, ["Metrik", "Değer"], bg=C_SUBHEAD)
    r += 1
    metrics = [
        ("Base NII (M)",         result.get("base_nii",   0)),
        ("NII at Risk (M)",      result.get("nii_at_risk", 0)),
        ("Total Assets (M)",     result.get("total_assets", 0)),
        ("Total Liabilities (M)",result.get("total_liabilities", 0)),
        ("Horizon (months)",     result.get("horizon_months", 12)),
    ]
    for i, (k, v) in enumerate(metrics):
        _write_data_row(ws1, r, [k, v], alt=(i % 2 == 0),
                        formats=[None, '#,##0.00'])
        r += 1

    r += 1
    # Scenario table
    scenarios = result.get("scenarios", {})
    if scenarios:
        _write_header_row(ws1, r, ["Senaryo", "NII (M)", "Δ NII (M)", "Δ NII (%)", "Yön"],
                          bg=C_NAVY)
        r += 1
        for i, (sc_name, sc) in enumerate(scenarios.items()):
            delta = sc.get("nii_change", 0)
            direction = "▲" if delta >= 0 else "▼"
            _write_data_row(ws1, r,
                [sc_name.replace("_", " ").title(),
                 sc.get("nii", 0),
                 delta,
                 sc.get("nii_change_pct", 0),
                 direction],
                alt=(i % 2 == 0),
                formats=[None, '#,##0.00', '#,##0.00', '0.00"%"', None])
            # Colour the Δ NII cell
            delta_cell = ws1.cell(row=r, column=3)
            delta_cell.font = Font(name="Calibri", size=9, bold=True,
                                   color=C_GREEN if delta >= 0 else C_RED)
            r += 1

    _set_col_width(ws1, "A", 28)
    _set_col_width(ws1, "B", 16)
    _set_col_width(ws1, "C", 16)
    _set_col_width(ws1, "D", 14)
    _set_col_width(ws1, "E", 10)

    # ── Sheet 2: Repricing Gap ────────────────────────────────────────────────
    ws2 = wb.create_sheet("Repricing Gap")
    _stamp_header(ws2, "Repricing Gap Table", "NII Sensitivity — Bucket Analysis")
    ws2.merge_cells("A1:G1")
    ws2.merge_cells("A2:G2")

    gap_table = result.get("repricing_gap", [])
    if gap_table:
        r2 = 4
        cols = ["Bucket", "RSA (M)", "RSL (M)", "Gap (M)", "Cum. Gap (M)",
                "Avg Asset Rate %", "Avg Liab Rate %"]
        _write_header_row(ws2, r2, cols, bg=C_NAVY)
        r2 += 1
        for i, row_data in enumerate(gap_table):
            gap = row_data.get("gap", 0)
            cum = row_data.get("cumulative_gap", 0)
            vals = [
                row_data.get("bucket", ""),
                row_data.get("rsa", 0),
                row_data.get("rsl", 0),
                gap,
                cum,
                row_data.get("avg_asset_rate", 0),
                row_data.get("avg_liability_rate", 0),
            ]
            fmts = [None, '#,##0.00', '#,##0.00', '#,##0.00', '#,##0.00', '0.00', '0.00']
            _write_data_row(ws2, r2, vals, alt=(i % 2 == 0), formats=fmts)
            # Colour gap cell
            gap_cell = ws2.cell(row=r2, column=4)
            gap_cell.font = Font(name="Calibri", size=9, bold=True,
                                 color=C_GREEN if gap >= 0 else C_RED)
            r2 += 1

        for col_idx, w in enumerate([14, 14, 14, 14, 14, 18, 18], start=1):
            ws2.column_dimensions[get_column_letter(col_idx)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ── EVE / DURATION GAP EXPORT ─────────────────────────────────────────────────

def export_eve(result: dict) -> io.BytesIO:
    """
    Export EVE / Duration Gap analysis to Excel.
    result = output of alm_engine.run_eve_analysis()
    """
    _require_openpyxl()
    wb = Workbook()

    # ── Sheet 1: EVE Summary ──────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "EVE Summary"
    _stamp_header(ws1, "EVE / Duration Gap Analysis", "Bek Rate Desk — IRRBB Module")
    ws1.merge_cells("A1:F1")
    ws1.merge_cells("A2:F2")

    r = 4
    _write_header_row(ws1, r, ["Metrik", "Değer", "Limit", "Durum"],
                      bg=C_SUBHEAD)
    r += 1

    base_eve = result.get("base_eve", 0)
    dur_gap  = result.get("duration_gap", 0)
    tier1    = result.get("tier1_capital", 1500)
    brsa_ok  = result.get("brsa_ok", True)
    worst_eve_chg = result.get("worst_eve_change", 0)
    worst_pct     = result.get("worst_eve_pct",    0)

    summary_rows = [
        ("Base EVE (M)",          base_eve,       "—",   "—"),
        ("Duration Gap (years)",  dur_gap,        "—",   "—"),
        ("Tier 1 Capital (M)",    tier1,          "—",   "—"),
        ("Worst ΔEVE (M)",        worst_eve_chg,  f"≤15% × {tier1:,.0f}M",
         "✓ OK" if brsa_ok else "✗ BREACH"),
        ("Worst ΔEVE (%T1)",      worst_pct,      "15%", "✓ OK" if brsa_ok else "✗ BREACH"),
    ]
    for i, (k, v, lim, st) in enumerate(summary_rows):
        _write_data_row(ws1, r, [k, v, lim, st], alt=(i % 2 == 0),
                        formats=[None, '#,##0.00', None, None])
        status_cell = ws1.cell(row=r, column=4)
        if st == "✓ OK":
            status_cell.font = Font(name="Calibri", size=9, bold=True, color=C_GREEN)
        elif st == "✗ BREACH":
            status_cell.font = Font(name="Calibri", size=9, bold=True, color=C_RED)
        r += 1

    r += 1
    # EVE Scenarios
    scenarios = result.get("scenarios", {})
    if scenarios:
        _write_header_row(ws1, r, ["Senaryo", "EVE (M)", "ΔEVE (M)", "ΔEVE (%T1)", "BRSA"],
                          bg=C_NAVY)
        r += 1
        for i, (sc_name, sc) in enumerate(scenarios.items()):
            delta = sc.get("eve_change", 0)
            pct   = sc.get("eve_change_pct_t1", 0)
            ok    = sc.get("brsa_ok", True)
            _write_data_row(ws1, r,
                [sc_name.replace("_", " ").title(),
                 sc.get("eve", 0), delta, pct,
                 "✓ OK" if ok else "✗ BREACH"],
                alt=(i % 2 == 0),
                formats=[None, '#,##0.00', '#,##0.00', '0.00"%"', None])
            ws1.cell(row=r, column=3).font = Font(name="Calibri", size=9, bold=True,
                                                   color=C_GREEN if delta >= 0 else C_RED)
            ws1.cell(row=r, column=5).font = Font(name="Calibri", size=9, bold=True,
                                                   color=C_GREEN if ok else C_RED)
            r += 1

    for col_idx, w in enumerate([28, 16, 16, 14, 12], start=1):
        ws1.column_dimensions[get_column_letter(col_idx)].width = w

    # ── Sheet 2: Asset Details ────────────────────────────────────────────────
    for sheet_label, details_key in [("Asset Details", "asset_details"),
                                      ("Liability Details", "liability_details")]:
        ws = wb.create_sheet(sheet_label)
        details = result.get(details_key, [])
        _stamp_header(ws, sheet_label, "EVE / Duration Gap Analysis")
        ws.merge_cells("A1:H1")
        r2 = 4
        cols2 = ["Name", "Amount (M)", "Rate %", "Maturity (y)",
                 "Type", "PV (M)", "Mod. Duration", "DV01"]
        _write_header_row(ws, r2, cols2, bg=C_NAVY)
        r2 += 1
        for i, item in enumerate(details):
            vals = [
                item.get("name", ""),
                item.get("amount", 0),
                item.get("rate", 0),
                item.get("maturity_years", 0),
                item.get("type", ""),
                item.get("pv", 0),
                item.get("mod_duration", 0),
                item.get("dv01", 0),
            ]
            fmts = [None, '#,##0.00', '0.00', '0.00', None, '#,##0.00', '0.0000', '#,##0.00']
            _write_data_row(ws, r2, vals, alt=(i % 2 == 0), formats=fmts)
            r2 += 1
        for col_idx, w in enumerate([22, 14, 10, 14, 10, 14, 14, 14], start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ── RV MATRIX EXPORT ──────────────────────────────────────────────────────────

def export_rv_matrix(result: dict) -> io.BytesIO:
    """
    Export RV Matrix analysis to Excel.
    result = output of dealer_analytics_engine.run_rv_matrix()
    """
    _require_openpyxl()
    wb = Workbook()

    ws = wb.active
    ws.title = "RV Matrix"
    _stamp_header(ws, "Relative Value (RV) Matrix", "Bek Rate Desk — Trading Desk")
    ws.merge_cells("A1:I1")
    ws.merge_cells("A2:I2")

    spreads = result.get("spreads", [])
    if not spreads:
        ws.cell(row=4, column=1, value="Veri bulunamadı")
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    r = 4
    cols = ["Spread", "Current (bps)", "Mean (bps)", "Std Dev",
            "Z-Score", "Min (1Y)", "Max (1Y)", "Signal", "Percentile"]
    _write_header_row(ws, r, cols, bg=C_NAVY)
    r += 1

    for i, sp in enumerate(spreads):
        z = sp.get("z_score")
        signal = ""
        if z is not None:
            if z > 2:   signal = "🔴 Çok Geniş"
            elif z > 1: signal = "🟡 Geniş"
            elif z < -2:signal = "🔵 Çok Dar"
            elif z < -1:signal = "🟢 Dar"
            else:        signal = "⚪ Nötr"

        vals = [
            sp.get("name", ""),
            sp.get("current_bps"),
            sp.get("mean_bps"),
            sp.get("std_bps"),
            round(z, 2) if z is not None else None,
            sp.get("min_1y_bps"),
            sp.get("max_1y_bps"),
            signal,
            sp.get("percentile"),
        ]
        fmts = [None, '0.0', '0.0', '0.0', '0.00', '0.0', '0.0', None, '0.0"%"']
        _write_data_row(ws, r, vals, alt=(i % 2 == 0), formats=fmts)

        # Colour z-score
        if z is not None:
            zc = ws.cell(row=r, column=5)
            zc.font = Font(name="Calibri", size=9, bold=True,
                           color=(C_RED if abs(z) > 2 else C_AMBER if abs(z) > 1 else C_GREEN))
        r += 1

    for col_idx, w in enumerate([18, 16, 14, 12, 12, 14, 14, 18, 14], start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ── REPRICING GAP STANDALONE EXPORT ──────────────────────────────────────────

def export_repricing_gap(result: dict) -> io.BytesIO:
    """
    Export standalone Repricing Gap table to Excel.
    result = output of alm_engine.run_nii_sensitivity()  (uses repricing_gap key)
    or any dict with a 'repricing_gap' list.
    """
    _require_openpyxl()
    wb = Workbook()
    ws = wb.active
    ws.title = "Repricing Gap"
    _stamp_header(ws, "Repricing Gap Analysis", "Bek Rate Desk — ALM Module")
    ws.merge_cells("A1:G1")
    ws.merge_cells("A2:G2")

    gap_table = result.get("repricing_gap", [])
    r = 4
    cols = ["Vade Dilimi", "RSA (M)", "RSL (M)", "Gap (M)",
            "Kümülatif Gap (M)", "Ort. Aktif Faiz %", "Ort. Pasif Faiz %"]
    _write_header_row(ws, r, cols, bg=C_NAVY)
    r += 1

    total_rsa, total_rsl = 0.0, 0.0
    for i, row_data in enumerate(gap_table):
        rsa  = row_data.get("rsa", 0)
        rsl  = row_data.get("rsl", 0)
        gap  = row_data.get("gap", 0)
        cum  = row_data.get("cumulative_gap", 0)
        total_rsa += rsa
        total_rsl += rsl
        vals = [
            row_data.get("bucket", ""),
            rsa, rsl, gap, cum,
            row_data.get("avg_asset_rate", 0),
            row_data.get("avg_liability_rate", 0),
        ]
        fmts = [None, '#,##0.00', '#,##0.00', '#,##0.00', '#,##0.00', '0.00', '0.00']
        _write_data_row(ws, r, vals, alt=(i % 2 == 0), formats=fmts)
        gap_cell = ws.cell(row=r, column=4)
        gap_cell.font = Font(name="Calibri", size=9, bold=True,
                             color=C_GREEN if gap >= 0 else C_RED)
        cum_cell = ws.cell(row=r, column=5)
        cum_cell.font = Font(name="Calibri", size=9, bold=True,
                             color=C_GREEN if cum >= 0 else C_RED)
        r += 1

    # Totals row
    _write_data_row(ws, r, ["TOPLAM", total_rsa, total_rsl, total_rsa - total_rsl, "", "", ""],
                    formats=[None, '#,##0.00', '#,##0.00', '#,##0.00', None, None, None])
    for ci in range(1, 8):
        ws.cell(row=r, column=ci).font = Font(name="Calibri", size=9, bold=True, color="000000")
        ws.cell(row=r, column=ci).fill = _fill("E2E8F0")

    for col_idx, w in enumerate([16, 14, 14, 14, 18, 20, 20], start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
